"""
Le public/MATRICULADOS2026.xlsx, atualiza matriculas 2026 (serie/turma) e
gera 12 cobrancas jan-dez 2026 com status=paga.

Mapping headers xlsx -> serie do banco:
  MATERNAL -> MATERNAL
  INFANTIL II / INFANTIL 2 / INFANTIL2 -> INFANTIL2
  INFANTIL III / INFANTIL 3 -> INFANTIL3
  INFANTIL IV / INFANTIL 4 / INFANTIL 5 -> INFANTIL4
  1o ANO ... 5o ANO -> 1º ANO ... 5º ANO
  6o ANO ... 9o ANO -> 6º ANO ... 9º ANO
  1a SERIE EM / 1a SERIE -> 1ª SÉRIE etc

Turma: A=MATUTINO, B=VESPERTINO. Header sem letra (so MATUTINO) -> MATUTINO.
"""
import re
import sys
import json
import unicodedata
import urllib.request
import urllib.error
from pathlib import Path

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

SUPABASE_URL = "http://127.0.0.1:55421"
SERVICE_KEY = "sb_secret_N7UND0UgjKTVK-Uodkm0Hg_xSvEMPvz"
ESCOLA_ID = "00000000-0000-0000-0000-000000000001"
PLANO_ID = "30000000-0000-0000-0000-000000000001"
ANO = 2026
XLSX = Path(__file__).parent.parent / "public" / "MATRICULADOS2026.xlsx"

HEADERS = {
    "apikey": SERVICE_KEY,
    "Authorization": f"Bearer {SERVICE_KEY}",
    "Content-Type": "application/json",
    "Prefer": "return=representation",
}

def api(method, path, body=None, params=""):
    url = f"{SUPABASE_URL}/rest/v1/{path}{params}"
    data = json.dumps(body).encode() if body is not None else None
    req = urllib.request.Request(url, data=data, headers=HEADERS, method=method)
    try:
        with urllib.request.urlopen(req) as r:
            raw = r.read()
            return json.loads(raw) if raw else []
    except urllib.error.HTTPError as e:
        msg = e.read().decode()
        raise RuntimeError(f"HTTP {e.code} {method} {path}: {msg}") from e

def norm(s):
    if s is None: return ""
    s = unicodedata.normalize("NFD", str(s))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    s = re.sub(r"\s+", " ", s).strip().upper()
    return s

# Mapping headers xlsx -> nome serie banco
SERIE_MAP = {
    "MATERNAL": "MATERNAL",
    "INFANTIL II": "INFANTIL2", "INFANTIL 2": "INFANTIL2", "INFANTIL2": "INFANTIL2",
    "INFANTIL III": "INFANTIL3", "INFANTIL 3": "INFANTIL3", "INFANTIL3": "INFANTIL3",
    "INFANTIL IV": "INFANTIL4", "INFANTIL 4": "INFANTIL4", "INFANTIL4": "INFANTIL4",
    "INFANTIL V": "INFANTIL4", "INFANTIL 5": "INFANTIL4",
    "1 ANO": "1º ANO", "2 ANO": "2º ANO", "3 ANO": "3º ANO", "4 ANO": "4º ANO", "5 ANO": "5º ANO",
    "6 ANO": "6º ANO", "7 ANO": "7º ANO", "8 ANO": "8º ANO", "9 ANO": "9º ANO",
    "1 SERIE": "1ª SÉRIE", "2 SERIE": "2ª SÉRIE", "3 SERIE": "3ª SÉRIE",
    "1 SERIE EM": "1ª SÉRIE", "2 SERIE EM": "2ª SÉRIE", "3 SERIE EM": "3ª SÉRIE",
}

def parse_header(text):
    """Parse 'MATERNAL - MATUTINO' or '1o ANO - A' or '1a SERIE - EM - A' -> (serie_nome, turno)."""
    t = norm(text).replace("º","").replace("ª","").replace("°","")
    t = re.sub(r"[ºª°]", "", t)
    parts = [p.strip() for p in t.split("-")]
    # tenta achar turno explicito (MATUTINO/VESPERTINO/NOTURNO/INTEGRAL)
    turno = None
    serie_label = None
    letra = None
    for p in parts:
        if p in ("MATUTINO","VESPERTINO","NOTURNO","INTEGRAL"):
            turno = p.lower()
        elif p == "EM":
            continue
        elif p in ("A","B","C","D"):
            letra = p
        else:
            serie_label = p
    if not serie_label:
        return None, None
    serie_nome = SERIE_MAP.get(serie_label)
    if not serie_nome:
        return None, None
    if turno is None:
        if letra == "B":
            turno = "vespertino"
        else:
            turno = "matutino"
    return serie_nome, turno

def parse_xlsx():
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    items = []  # list of {nome, valor, serie_nome, turno}
    current_serie = None
    current_turno = None
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        for row in ws.iter_rows(values_only=True):
            # localiza header em qualquer coluna (texto contendo " - " ou ANO/SERIE)
            for cell in row:
                if not cell or not isinstance(cell, str): continue
                up = norm(cell)
                if (" - " in up) and ("ANO" in up or "SERIE" in up or "MATERNAL" in up or "INFANTIL" in up):
                    s, t = parse_header(cell)
                    if s:
                        current_serie = s
                        current_turno = t
                    break
            # tenta extrair linha de aluno: col3=#, col4=nome, col5=valor (estrutura padrao)
            if current_serie is None: continue
            # encontra primeira string longa = nome e primeiro num = valor
            nome = None
            valor = None
            for cell in row:
                if isinstance(cell, str) and len(cell.strip()) > 4 and not any(k in norm(cell) for k in ("ALUNO","MATRICULA","MENSALIDADE"," - ")):
                    nome = cell.strip()
                    break
            for cell in row:
                if isinstance(cell, (int, float)) and 100 <= cell <= 5000:
                    valor = float(cell)
                    break
            if nome and valor:
                items.append({"nome": nome, "valor": valor, "serie_nome": current_serie, "turno": current_turno})
    return items

def get_series_cache():
    rows = api("GET", "series", params=f"?escola_id=eq.{ESCOLA_ID}&select=id,nome")
    return {r["nome"]: r["id"] for r in rows}

def get_or_create_turma(serie_id, turno):
    """Turma 2026 com nome = turno uppercase."""
    nome = turno.upper()
    rows = api("GET", "turmas", params=f"?escola_id=eq.{ESCOLA_ID}&serie_id=eq.{serie_id}&ano_letivo=eq.{ANO}&turno=eq.{turno}&nome=eq.{nome}&select=id")
    if rows:
        return rows[0]["id"]
    created = api("POST", "turmas", body=[{
        "escola_id": ESCOLA_ID,
        "serie_id": serie_id,
        "nome": nome,
        "ano_letivo": ANO,
        "turno": turno,
        "ativo": True,
    }])
    return created[0]["id"]

def find_aluno(nome):
    # match exato normalizado
    rows = api("GET", "alunos", params=f"?escola_id=eq.{ESCOLA_ID}&select=id,nome")
    target = norm(nome)
    for r in rows:
        if norm(r["nome"]) == target:
            return r["id"]
    return None

def upsert_matricula(aluno_id, serie_id, turma_id):
    rows = api("GET", "matriculas", params=f"?escola_id=eq.{ESCOLA_ID}&aluno_id=eq.{aluno_id}&ano_letivo=eq.{ANO}&select=id")
    if rows:
        mid = rows[0]["id"]
        api("PATCH", "matriculas", body={"serie_id": serie_id, "turma_id": turma_id, "status": "ativa"},
            params=f"?id=eq.{mid}")
        return mid
    # cria
    codigo = f"M{ANO}-{aluno_id[:8]}"
    created = api("POST", "matriculas", body=[{
        "escola_id": ESCOLA_ID,
        "aluno_id": aluno_id,
        "serie_id": serie_id,
        "turma_id": turma_id,
        "plano_id": PLANO_ID,
        "codigo": codigo,
        "data_matricula": f"{ANO}-01-01",
        "ano_letivo": ANO,
        "status": "ativa",
    }])
    return created[0]["id"]

def gen_cobrancas(matricula_id, aluno_id, valor):
    # apaga existentes 2026 desta matricula
    api("DELETE", "cobrancas", params=f"?matricula_id=eq.{matricula_id}&competencia=like.{ANO}-*")
    rows = []
    for m in range(1, 13):
        comp = f"{ANO}-{m:02d}"
        venc = f"{ANO}-{m:02d}-10"
        rows.append({
            "escola_id": ESCOLA_ID,
            "aluno_id": aluno_id,
            "matricula_id": matricula_id,
            "plano_id": PLANO_ID,
            "descricao": f"Mensalidade {m:02d}/{ANO}",
            "competencia": comp,
            "numero_parcela": m,
            "valor_original": valor,
            "valor_desconto": 0,
            "valor_acrescimo": 0,
            "data_vencimento": venc,
            "status": "paga",
        })
    cobrancas = api("POST", "cobrancas", body=rows)
    # pagamento para cada
    pags = [{
        "escola_id": ESCOLA_ID,
        "cobranca_id": c["id"],
        "aluno_id": aluno_id,
        "matricula_id": matricula_id,
        "valor_pago": valor,
        "data_pagamento": c["data_vencimento"],
        "forma_pagamento": "transferencia",
        "observacao": "Financeiro terceirizado",
    } for c in cobrancas]
    api("POST", "pagamentos", body=pags)
    return len(cobrancas)

def main():
    if not XLSX.exists():
        print(f"Arquivo nao encontrado: {XLSX}"); sys.exit(1)

    print(f"Parseando {XLSX.name}...")
    items = parse_xlsx()
    print(f"  {len(items)} linhas de aluno extraidas")

    series_cache = get_series_cache()
    turmas_cache = {}

    updated_mat = 0
    created_mat = 0
    cobrancas_created = 0
    skipped = []
    errors = []

    for item in items:
        try:
            serie_id = series_cache.get(item["serie_nome"])
            if not serie_id:
                skipped.append((item["nome"], f"serie nao mapeada: {item['serie_nome']}"))
                continue
            tkey = (serie_id, item["turno"])
            if tkey not in turmas_cache:
                turmas_cache[tkey] = get_or_create_turma(serie_id, item["turno"])
            turma_id = turmas_cache[tkey]

            aluno_id = find_aluno(item["nome"])
            if not aluno_id:
                skipped.append((item["nome"], "aluno nao encontrado"))
                continue

            existing = api("GET", "matriculas", params=f"?escola_id=eq.{ESCOLA_ID}&aluno_id=eq.{aluno_id}&ano_letivo=eq.{ANO}&select=id")
            mat_id = upsert_matricula(aluno_id, serie_id, turma_id)
            if existing: updated_mat += 1
            else: created_mat += 1

            n = gen_cobrancas(mat_id, aluno_id, item["valor"])
            cobrancas_created += n
        except Exception as e:
            errors.append((item["nome"], str(e)[:200]))

    print(f"\n=== Resultado ===")
    print(f"Matriculas atualizadas: {updated_mat}")
    print(f"Matriculas criadas:     {created_mat}")
    print(f"Cobrancas geradas:      {cobrancas_created}")
    print(f"Pulados:                {len(skipped)}")
    print(f"Erros:                  {len(errors)}")
    if skipped[:10]:
        print(f"\nPrimeiros pulados:")
        for nome, motivo in skipped[:10]:
            print(f"  - {nome}: {motivo}")
    if errors[:10]:
        print(f"\nPrimeiros erros:")
        for nome, motivo in errors[:10]:
            print(f"  - {nome}: {motivo}")

if __name__ == "__main__":
    main()
